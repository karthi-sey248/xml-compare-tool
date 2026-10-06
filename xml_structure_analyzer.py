#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
xml_structure_analyzer.py
--------------------------
Oru XML file eduthu, athula Front Matter / Body Matter / Back Matter
(and anything else, "Other") nu section-a pirichi, ovoru section-layum
enna enna TAG use pannirukanga, enna CLASS use pannirukanga, enna
ATTRIBUTES irukanga nu ella-athayum kandu pidichu, oru colored Excel
report generate pannura script.

Section detection eppadi work aagudhu (generic, any XML):
    - Document ah top-to-bottom recursive-a walk pannum.
    - Edhavadhu element ku class="front" kandaal, andha element and
      adhoda ELLA children-um "Front Matter" section-a mark aagum.
    - class="body" -> "Body Matter". class="back" -> "Back Matter".
    - Idhu edhuvum kidaikama irundha, andha element "Other / Root Level"
      section-a pogum (e.g. <item-info>, <jid>, <doi> mari root-level
      metadata tags).

USAGE:
    1. Keela INPUT / OUTPUT SETTINGS la, unga XML file path & output
       excel path mention pannunga.
    2. python3 xml_structure_analyzer.py
"""

import xml.etree.ElementTree as ET
import xml.parsers.expat as expat
from collections import OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ============================================================
#                 INPUT / OUTPUT SETTINGS
# ============================================================
INPUT_XML_PATH = "/mnt/user-data/uploads/tx2.xhtml"
OUTPUT_EXCEL_PATH = "/mnt/user-data/outputs/xml_structure_report.xlsx"
# ============================================================


SECTION_ORDER = ["Front Matter", "Body Matter", "Back Matter", "Other / Root Level"]
SECTION_COLORS = {
    "Front Matter": "DDEBF7",        # light blue
    "Body Matter": "D9F2D9",         # light green
    "Back Matter": "FDE9D9",         # light orange
    "Other / Root Level": "EDEDED",  # light grey
}
SECTION_HEADER_COLORS = {
    "Front Matter": "2E74B5",
    "Body Matter": "38761D",
    "Back Matter": "B45F06",
    "Other / Root Level": "666666",
}


def parse_with_line_numbers(path):
    """XML file ah parse pannitu, ovoru element ku athoda start LINE
    NUMBER um track pannum."""
    tb = ET.TreeBuilder()
    line_map = {}

    parser = expat.ParserCreate()
    parser.buffer_text = True

    def start(tag, attrs):
        el = tb.start(tag, attrs)
        line_map[id(el)] = parser.CurrentLineNumber

    def end(tag):
        tb.end(tag)

    def data(text):
        tb.data(text)

    parser.StartElementHandler = start
    parser.EndElementHandler = end
    parser.CharacterDataHandler = data

    with open(path, "rb") as f:
        parser.ParseFile(f)

    root = tb.close()
    return root, line_map


def local_tag(tag):
    return tag.split("}", 1)[1] if "}" in tag else tag


def norm(text, max_len=70):
    if not text:
        return ""
    t = " ".join(text.split())
    if len(t) > max_len:
        t = t[:max_len] + "..."
    return t


def inner_text_sample(elem, max_len=90):
    """Element ku direct text illama, ella nested children content-um
    (combined) sample-a return pannum, so container tags kum meaningful
    sample text kaanum."""
    parts = []

    def rec(e):
        if e.text:
            parts.append(e.text)
        for c in e:
            rec(c)
            if c.tail:
                parts.append(c.tail)

    rec(elem)
    return norm(" ".join(parts), max_len)


def classify_section(elem):
    cls = (elem.get("class") or "").strip().lower()
    if cls == "front":
        return "Front Matter"
    if cls == "body":
        return "Body Matter"
    if cls == "back":
        return "Back Matter"
    return None


class SectionStats:
    def __init__(self):
        # key: (section, tag, class) -> dict
        self.data = OrderedDict()

    def record(self, section, elem, line_map):
        tag = local_tag(elem.tag)
        cls = elem.get("class") or ""
        key = (section, tag, cls)
        if key not in self.data:
            self.data[key] = {
                "section": section,
                "tag": tag,
                "class": cls,
                "attr_names": OrderedDict(),  # attr name -> example value
                "count": 0,
                "example_line": line_map.get(id(elem)),
                "example_text": inner_text_sample(elem),
            }
        entry = self.data[key]
        entry["count"] += 1
        for k, v in elem.attrib.items():
            if k == "class":
                continue
            if k not in entry["attr_names"]:
                entry["attr_names"][k] = v


def walk(elem, current_section, stats, line_map):
    new_section = classify_section(elem) or current_section
    stats.record(new_section, elem, line_map)
    for child in elem:
        walk(child, new_section, stats, line_map)


def analyze(xml_path):
    root, line_map = parse_with_line_numbers(xml_path)
    stats = SectionStats()
    walk(root, "Other / Root Level", stats, line_map)
    return stats


def write_excel(stats, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "XML Structure Report"

    headers = ["S.No", "Section", "Tag", "Class", "Attributes Used",
               "Occurrence Count", "First Line No", "Sample Text"]
    ws.append(headers)
    header_font = Font(color="FFFFFF", bold=True)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = PatternFill(start_color="1F2A44", end_color="1F2A44", fill_type="solid")
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # sort rows: by section order, then tag name, then class name
    def sort_key(item):
        (section, tag, cls) = item[0]
        sec_idx = SECTION_ORDER.index(section) if section in SECTION_ORDER else 99
        return (sec_idx, tag, cls)

    rows = sorted(stats.data.items(), key=sort_key)

    sno = 1
    for (section, tag, cls), entry in rows:
        attrs_str = ", ".join(f"{k}" for k in entry["attr_names"].keys())
        ws.append([
            sno, section, tag, cls, attrs_str, entry["count"],
            entry["example_line"], entry["example_text"]
        ])
        fill_color = SECTION_COLORS.get(section, "FFFFFF")
        for col in range(1, len(headers) + 1):
            ws.cell(row=sno + 1, column=col).fill = PatternFill(
                start_color=fill_color, end_color=fill_color, fill_type="solid"
            )
        ws.cell(row=sno + 1, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=sno + 1, column=6).alignment = Alignment(horizontal="center")
        ws.cell(row=sno + 1, column=7).alignment = Alignment(horizontal="center")
        sno += 1

    widths = [8, 20, 22, 22, 45, 16, 14, 45]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{sno}"

    # ---- Summary sheet ----
    summary = wb.create_sheet("Summary")
    summary.append(["Section", "Distinct Tag/Class Combinations", "Total Elements"])
    for col in range(1, 4):
        cell = summary.cell(row=1, column=col)
        cell.fill = PatternFill(start_color="1F2A44", end_color="1F2A44", fill_type="solid")
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    section_totals = OrderedDict((s, [0, 0]) for s in SECTION_ORDER)
    for (section, tag, cls), entry in stats.data.items():
        if section not in section_totals:
            section_totals[section] = [0, 0]
        section_totals[section][0] += 1
        section_totals[section][1] += entry["count"]

    r = 2
    for section, (distinct, total) in section_totals.items():
        summary.append([section, distinct, total])
        color = SECTION_COLORS.get(section, "FFFFFF")
        for col in range(1, 4):
            summary.cell(row=r, column=col).fill = PatternFill(
                start_color=color, end_color=color, fill_type="solid"
            )
        r += 1

    for i, w in enumerate([26, 30, 16], start=1):
        summary.column_dimensions[get_column_letter(i)].width = w

    wb.save(output_path)


if __name__ == "__main__":
    print(f"Analyzing: {INPUT_XML_PATH}")
    stats = analyze(INPUT_XML_PATH)
    write_excel(stats, OUTPUT_EXCEL_PATH)
    print(f"Found {len(stats.data)} distinct tag/class combinations.")
    print(f"Report saved to: {OUTPUT_EXCEL_PATH}")
